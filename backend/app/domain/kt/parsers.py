"""Bounded document parsers producing anchored extracted text.

Parsers only READ: nothing in a document is executed, rendered or followed.
What a parser cannot read (tables, images, scanned pages, active content) is
reported as an `ExtractionWarning` instead of being dropped silently - OCR and
table/image extraction need a qualified parser (ADR 0020) and are not done here.

`run_with_timeout` bounds a parse in-process (a best-effort thread bound that
cannot kill a stuck native call). The ingestion worker uses the subprocess
runner in `parse_isolation.py`, which can: use that for anything uploaded.
"""

from __future__ import annotations

import re
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeoutError
from html.parser import HTMLParser
from typing import Optional, Protocol
from uuid import uuid4

from app.domain.kt.limits import MAX_EXTRACTED_TEXT_BYTES, MAX_PAGES, PARSE_TIMEOUT_SECONDS
from app.domain.kt.models import ExtractedAnchor, ExtractionWarning

_parse_pool = ThreadPoolExecutor(max_workers=4, thread_name_prefix="kt-parse")

ANCHOR_TEXT_LIMIT = 4000  # matches ExtractedAnchor.text; longer paragraphs are split, never cut
MAX_WARNING_PAGES = 200


class ParseTimeoutError(Exception):
    pass


class ParseRejectedError(Exception):
    """The document is understood to be unusable; `code` is stable, `message` actionable."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class ParseResult:
    def __init__(
        self,
        anchors: list[ExtractedAnchor],
        total_pages: int | None,
        truncated: bool,
        warnings: Optional[list[ExtractionWarning]] = None,
        stats: Optional[dict] = None,
    ) -> None:
        self.anchors = anchors
        self.total_pages = total_pages
        self.truncated = truncated
        self.warnings = warnings or []
        self.stats = stats or {}


class DocumentParser(Protocol):
    def parse(self, content: bytes) -> ParseResult: ...


def run_with_timeout(fn, *args, timeout: float = PARSE_TIMEOUT_SECONDS):
    future = _parse_pool.submit(fn, *args)
    try:
        return future.result(timeout=timeout)
    except FutureTimeoutError as exc:
        raise ParseTimeoutError(f"parse job exceeded {timeout}s") from exc


def _split_text(text: str, limit: int = ANCHOR_TEXT_LIMIT) -> list[str]:
    """Split on whitespace boundaries so no anchor silently loses its tail."""
    text = text.strip()
    if len(text) <= limit:
        return [text] if text else []
    pieces: list[str] = []
    while len(text) > limit:
        cut = text.rfind(" ", 0, limit)
        cut = cut if cut > limit // 2 else limit
        pieces.append(text[:cut].strip())
        text = text[cut:].strip()
    if text:
        pieces.append(text)
    return [p for p in pieces if p]


def _make_anchor(page: int | None, paragraph_index: int, text: str, section: Optional[str] = None) -> ExtractedAnchor:
    return ExtractedAnchor(
        anchor_id=uuid4().hex, page=page, paragraph_index=paragraph_index, text=text[:ANCHOR_TEXT_LIMIT],
        section=(section[:200] if section else None),
    )


class _Collector:
    """Accumulates anchors under the shared text-size cap."""

    def __init__(self) -> None:
        self.anchors: list[ExtractedAnchor] = []
        self.total_bytes = 0
        self.truncated = False

    def add(self, page: Optional[int], raw: str, section: Optional[str] = None) -> bool:
        """False once the cap is hit (callers stop reading)."""
        for piece in _split_text(raw):
            self.total_bytes += len(piece.encode("utf-8"))
            if self.total_bytes > MAX_EXTRACTED_TEXT_BYTES:
                self.truncated = True
                return False
            self.anchors.append(_make_anchor(page, len(self.anchors), piece, section))
        return True


def _truncation_warning(truncated: bool) -> list[ExtractionWarning]:
    if not truncated:
        return []
    return [ExtractionWarning(code="truncated", message="The document is longer than the extraction limit; the end was not read.")]


class TxtParser:
    def parse(self, content: bytes) -> ParseResult:
        text = content.decode("utf-8", errors="replace")
        collector = _Collector()
        for para in text.split("\n\n"):
            if not collector.add(None, para):
                break
        return ParseResult(collector.anchors, None, collector.truncated, _truncation_warning(collector.truncated))


_HIDDEN_STYLE = re.compile(r"display\s*:\s*none|visibility\s*:\s*hidden", re.IGNORECASE)
_BLOCK_TAGS = {
    "p", "div", "li", "ul", "ol", "tr", "br", "section", "article", "aside", "header", "footer", "main", "nav",
    "blockquote", "pre", "table", "dd", "dt", "dl", "figure", "figcaption", "h1", "h2", "h3", "h4", "h5", "h6",
}
_SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "head", "iframe", "object", "embed"}
_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_VOID_TAGS = {"br", "img", "hr", "input", "meta", "link", "area", "base", "col", "embed", "source", "track", "wbr"}


class _TextExtractor(HTMLParser):
    """Visible text only. Scripts, styles, comments and elements hidden with the
    `hidden` attribute or inline CSS are dropped - hidden text is a classic way to
    smuggle instructions to a model that reads page content."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[Optional[str], str]] = []  # (section, text)
        self._buffer: list[str] = []
        self._skip_depth = 0
        self._hidden_depth = 0
        self._heading: Optional[str] = None
        self._in_heading: Optional[str] = None
        self.section: Optional[str] = None
        self.tables = 0
        self.images = 0
        self.hidden_elements = 0

    def _flush(self) -> None:
        text = " ".join("".join(self._buffer).split())
        self._buffer = []
        if not text:
            return
        if self._in_heading:
            self.section = text[:200]
        self.blocks.append((self.section, text))

    def handle_starttag(self, tag, attrs):
        attr = dict(attrs)
        hidden = "hidden" in attr or bool(_HIDDEN_STYLE.search(attr.get("style") or ""))
        if tag in _VOID_TAGS:
            if tag == "img":
                self.images += 1
            if tag in _BLOCK_TAGS and not (self._skip_depth or self._hidden_depth):
                self._flush()
            return
        if tag in _SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if hidden or self._hidden_depth:
            if hidden:
                self.hidden_elements += 1
            self._hidden_depth += 1
            return
        if tag == "table":
            self.tables += 1
        if tag in _BLOCK_TAGS:
            self._flush()
        if tag in _HEADINGS:
            self._in_heading = tag

    def handle_endtag(self, tag):
        if tag in _VOID_TAGS:
            return
        if tag in _SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if self._hidden_depth:
            self._hidden_depth -= 1
            return
        if tag in _BLOCK_TAGS:
            self._flush()
        if tag in _HEADINGS:
            self._in_heading = None

    def handle_data(self, data):
        if not (self._skip_depth or self._hidden_depth):
            self._buffer.append(data)

    def close(self):
        super().close()
        self._flush()


class HtmlParser:
    def parse(self, content: bytes) -> ParseResult:
        extractor = _TextExtractor()
        extractor.feed(content.decode("utf-8", errors="replace"))
        extractor.close()

        collector = _Collector()
        for section, text in extractor.blocks:
            if not collector.add(None, text, section):
                break

        warnings = _truncation_warning(collector.truncated)
        if extractor.tables:
            warnings.append(ExtractionWarning(
                code="tables_not_extracted", count=extractor.tables,
                message=f"{extractor.tables} table(s) were read as plain text only; row/column structure is not preserved.",
            ))
        if extractor.images:
            warnings.append(ExtractionWarning(
                code="images_not_extracted", count=extractor.images,
                message=f"{extractor.images} image(s) were not read (no OCR); text inside images is missing.",
            ))
        if extractor.hidden_elements:
            warnings.append(ExtractionWarning(
                code="hidden_content_skipped", count=extractor.hidden_elements,
                message=f"{extractor.hidden_elements} hidden element(s) were ignored.",
            ))
        stats = {"tables": extractor.tables, "images": extractor.images, "hidden_elements": extractor.hidden_elements}
        return ParseResult(collector.anchors, None, collector.truncated, warnings, stats)


class DocxParser:
    def parse(self, content: bytes) -> ParseResult:
        import io
        import zipfile

        import docx  # python-docx

        document = docx.Document(io.BytesIO(content))
        collector = _Collector()
        section: Optional[str] = None
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                section = text[:200]
            if not collector.add(None, text, section):
                break

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            images = sum(1 for name in archive.namelist() if name.startswith("word/media/"))
        tables = len(document.tables)

        warnings = _truncation_warning(collector.truncated)
        if tables:
            warnings.append(ExtractionWarning(
                code="tables_not_extracted", count=tables,
                message=f"{tables} table(s) were not extracted; their text is missing from the source.",
            ))
        if images:
            warnings.append(ExtractionWarning(
                code="images_not_extracted", count=images,
                message=f"{images} embedded image(s) were not read (no OCR); text inside images is missing.",
            ))
        return ParseResult(collector.anchors, None, collector.truncated, warnings, {"tables": tables, "images": images})


_PDF_ACTIVE_MARKERS = (b"/JavaScript", b"/JS ", b"/Launch", b"/EmbeddedFile", b"/OpenAction")


class PdfParser:
    def parse(self, content: bytes) -> ParseResult:
        import io

        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ParseRejectedError("encrypted", "The PDF is password-protected. Remove the password and upload it again.")

        collector = _Collector()
        total_pages = len(reader.pages)
        pages_to_read = min(total_pages, MAX_PAGES)
        truncated = total_pages > MAX_PAGES

        empty_pages: list[int] = []
        image_count = 0
        replacement_chars = 0
        text_chars = 0
        for page_number in range(pages_to_read):
            page = reader.pages[page_number]
            page_text = page.extract_text() or ""
            try:
                image_count += len(page.images)
            except Exception:  # noqa: BLE001 - image enumeration is advisory only
                pass
            if not page_text.strip():
                empty_pages.append(page_number + 1)
                continue
            text_chars += len(page_text)
            replacement_chars += page_text.count("\ufffd")
            # pypdf separates lines with a single newline, so paragraphs are split on blank lines only.
            for para in re.split(r"\n\s*\n", page_text):
                if not collector.add(page_number + 1, para):
                    truncated = True
                    break
            if collector.truncated:
                truncated = True
                break

        warnings = _truncation_warning(truncated)
        if empty_pages and len(empty_pages) < pages_to_read:
            warnings.append(ExtractionWarning(
                code="pages_without_text", count=len(empty_pages), pages=empty_pages[:MAX_WARNING_PAGES],
                message=f"{len(empty_pages)} page(s) have no extractable text (scanned or image-only); they were not read.",
            ))
        if image_count:
            warnings.append(ExtractionWarning(
                code="images_not_extracted", count=image_count,
                message=f"{image_count} image(s) were not read (no OCR); text inside images is missing.",
            ))
        if text_chars and replacement_chars / text_chars > 0.05:
            warnings.append(ExtractionWarning(
                code="low_confidence_text", count=replacement_chars,
                message="A large share of the text could not be decoded (custom fonts or damaged encoding); review the extract before relying on it.",
            ))
        if any(marker in content for marker in _PDF_ACTIVE_MARKERS):
            warnings.append(ExtractionWarning(
                code="active_content_ignored",
                message="The PDF contains scripts, launch actions or embedded files; they were ignored.",
            ))
        if pages_to_read and len(empty_pages) == pages_to_read:
            raise ParseRejectedError(
                "no_text",
                "The PDF has no extractable text - it looks scanned or image-only. OCR is not available; upload a text-based PDF or DOCX.",
            )
        stats = {"pages": total_pages, "empty_pages": len(empty_pages), "images": image_count}
        return ParseResult(collector.anchors, total_pages, truncated, warnings, stats)


PARSERS_BY_KIND: dict[str, DocumentParser] = {
    "txt": TxtParser(),
    "html": HtmlParser(),
    "docx": DocxParser(),
    "pdf": PdfParser(),
}

PARSER_LABELS = {"txt": "plain-text", "html": "html.parser", "docx": "python-docx", "pdf": "pypdf"}


def parse_bounded(kind: str, content: bytes) -> ParseResult:
    parser = PARSERS_BY_KIND.get(kind)
    if parser is None:
        raise ValueError(f"no parser registered for kind={kind!r}")
    return run_with_timeout(parser.parse, content)


def require_text(result: ParseResult, kind: str) -> ParseResult:
    if not result.anchors:
        raise ParseRejectedError(
            "no_text", f"No readable text was found in this {kind.upper()} document. Check that it is not empty or image-only."
        )
    return result
