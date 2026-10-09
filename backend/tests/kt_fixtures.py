"""Synthetic documents for the KT ingestion tests (Prompt 4, ADR 0020). Test support, not collected."""

from __future__ import annotations

import base64
import io
from typing import Optional

_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def make_pdf(pages: list[Optional[str]]) -> bytes:
    """A small valid PDF. A page given as None has no text (what a scanned page looks like to a text extractor)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    add(b"<< /Type /Catalog /Pages 2 0 R >>")  # 1
    kids_placeholder = len(objects) + 1  # 2: patched below once the page ids are known
    add(b"")
    add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")  # 3
    page_ids = []
    for text in pages:
        if text is None:
            content_id = add(b"<< /Length 0 >>\nstream\n\nendstream")
        else:
            safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1", "replace")
            content_id = add(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        page_ids.append(add(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        ))
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects[kids_placeholder - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_encrypted_pdf(password: str = "secret") -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["secret text"]))))
    writer.encrypt(password, algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(
    paragraphs: list[str] | None = None, *, headings: list[str] | None = None, tables: int = 0, images: int = 0
) -> bytes:
    import docx

    document = docx.Document()
    for heading in headings or []:
        document.add_heading(heading, level=1)
    for text in paragraphs or []:
        document.add_paragraph(text)
    for _ in range(tables):
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "cell text that is not extracted"
    for _ in range(images):
        document.add_picture(io.BytesIO(_TINY_PNG))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def inproc_parse(kind: str, content: bytes):
    """Production parsers without the child process (the child itself is covered separately)."""
    from app.domain.kt.parse_isolation import IngestionParseError
    from app.domain.kt.parsers import ParseRejectedError, parse_bounded, require_text

    try:
        return require_text(parse_bounded(kind, content), kind)
    except ParseRejectedError as exc:
        raise IngestionParseError(exc.code, exc.message) from exc
    except Exception as exc:  # noqa: BLE001 - the child process maps these to "corrupt"
        raise IngestionParseError("corrupt", f"could not be read ({type(exc).__name__})") from exc
